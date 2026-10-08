"""
Script to compile PROJECT_REPORT.md into a professional Microsoft Word (.docx) report
with executive styling, shaded tables, callout blocks, and clean typography.
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PROJECT_ROOT = Path(__file__).resolve().parent.parent
MD_PATH = PROJECT_ROOT / "PROJECT_REPORT.md"
OUTPUT_DOCX_REPO = PROJECT_ROOT / "MLOps_Project_Report.docx"
OUTPUT_DOCX_DOWNLOADS = Path("c:/Users/himan/Downloads/MLOps_Project_Report.docx")

# Color Palette
COLOR_NAVY = RGBColor(30, 58, 138)       # #1E3A8A - Primary Header
COLOR_ROYAL = RGBColor(37, 99, 235)      # #2563EB - Secondary Header
COLOR_SLATE = RGBColor(51, 65, 85)       # #334155 - Subtitle / Dark text
COLOR_BODY = RGBColor(15, 23, 42)        # #0F172A - Body text
COLOR_MUTED = RGBColor(100, 116, 139)    # #64748B - Muted captions
HEX_TABLE_HEADER = "1E3A8A"
HEX_ZEBRA = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_CALLOUT_BG = "F1F5F9"
HEX_CALLOUT_BORDER = "2563EB"


def set_cell_background(cell, fill_hex):
    """Sets background shading of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding of a table cell (values in dxa: 20 dxa = 1 pt)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def add_callout(doc, text, title="NOTE / ARCHITECTURE"):
    """Adds an executive styled callout box with a left vertical border."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, HEX_CALLOUT_BG)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    # Set left border thick, others none
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{HEX_CALLOUT_BORDER}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(tc_borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(f"■ {title}\n")
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(10)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_ROYAL

    run_body = p.add_run(text)
    run_body.font.name = "Calibri"
    run_body.font.size = Pt(9.5)
    run_body.font.color.rgb = COLOR_BODY

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_styled_paragraph(doc, text, style='Normal', space_after=6, bold=False, italic=False, color=COLOR_BODY, size=11):
    """Adds a paragraph with inline markdown parsing (**bold**, `code`, etc.)."""
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15

    # Parse inline bold and code
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)')
    tokens = pattern.split(text)

    for token in tokens:
        if not token:
            continue
        run = p.add_run()
        run.font.name = 'Calibri'
        run.font.size = Pt(size)
        run.font.color.rgb = color

        if token.startswith('**') and token.endswith('**'):
            run.text = token[2:-2]
            run.font.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run.text = token[1:-1]
            run.font.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run.text = token[1:-1]
            run.font.name = 'Consolas'
            run.font.size = Pt(size - 0.5)
            run.font.color.rgb = RGBColor(194, 65, 12)  # Rust/Amber code color
        else:
            run.text = token
            run.font.bold = bold
            run.font.italic = italic

    return p


def main():
    if not MD_PATH.exists():
        raise FileNotFoundError(f"Source markdown not found: {MD_PATH}")

    content = MD_PATH.read_text(encoding="utf-8")
    doc = Document()

    # Set page layout to 1-inch margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # -------------------------------------------------------------
    # Formal Academic Cover Page / Header
    # -------------------------------------------------------------
    inst_p = doc.add_paragraph()
    inst_p.paragraph_format.space_before = Pt(8)
    inst_p.paragraph_format.space_after = Pt(2)
    inst_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = inst_p.add_run("APEX INSTITUTE OF TECHNOLOGY (AIT)\nDEPARTMENT OF COMPUTER SCIENCE & ENGINEERING")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(13)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_ROYAL

    course_p = doc.add_paragraph()
    course_p.paragraph_format.space_before = Pt(4)
    course_p.paragraph_format.space_after = Pt(22)
    course_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_course = course_p.add_run("DATA ENGINEERING II ― CAPSTONE PROJECT REPORT")
    r_course.font.name = "Calibri"
    r_course.font.size = Pt(11)
    r_course.font.bold = True
    r_course.font.color.rgb = COLOR_SLATE

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(8)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("End-to-End MLOps Pipeline: Automated Model Deployment, Containerization with Kubernetes, & Real-Time Observability")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_NAVY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(24)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("Credit Default Risk Prediction • Kolmogorov-Smirnov & PSI Drift Engine • Prometheus & Grafana Telemetry")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11.5)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_SLATE

    # Academic Credentials Box (Table)
    cred_table = doc.add_table(rows=6, cols=2)
    cred_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cred_data = [
        ("Student Name", "HIMANSHU GOYAL"),
        ("UID", "24BDA70369"),
        ("Class / Section", "24BDS 4NTPP"),
        ("Course / Subject", "Data Engineering II"),
        ("Department", "AIT - CSE Department"),
        ("Faculty Evaluator / Mentor", "Mr. Deepak Kumar"),
    ]

    for row_idx, (label, val) in enumerate(cred_data):
        c0 = cred_table.cell(row_idx, 0)
        c1 = cred_table.cell(row_idx, 1)
        c0.width = Inches(2.5)
        c1.width = Inches(4.0)

        set_cell_background(c0, "1E3A8A" if row_idx == 0 else "F1F5F9")
        set_cell_background(c1, "2563EB" if row_idx == 0 else "FFFFFF")
        set_cell_margins(c0, top=90, bottom=90, left=140, right=140)
        set_cell_margins(c1, top=90, bottom=90, left=140, right=140)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(label)
        r0.font.name = "Calibri"
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(255, 255, 255) if row_idx == 0 else COLOR_NAVY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(255, 255, 255) if row_idx == 0 else COLOR_BODY

    # Space and Page Break for clean formal start
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    doc.add_page_break()


    # Parse lines of markdown
    lines = content.splitlines()
    i = 0
    in_mermaid = False
    mermaid_buffer = []

    while i < len(lines):
        line = lines[i].strip()

        # Skip main markdown H1/H2 header that we already styled above
        if line.startswith("# Comprehensive Project Report") or line.startswith("## End-to-End MLOps Pipeline"):
            i += 1
            continue

        # Skip horizontal dividers
        if line in ("---", "***"):
            i += 1
            continue

        # Handle Mermaid Code Blocks
        if line.startswith("```mermaid"):
            in_mermaid = True
            mermaid_buffer = []
            i += 1
            continue
        elif in_mermaid:
            if line.startswith("```"):
                in_mermaid = False
                diagram_summary = (
                    "Microservice Architecture Flow:\n"
                    "1. Developer Git Push triggers GitHub Actions 5-stage automated CI/CD pipeline.\n"
                    "2. Model Quality Gate enforces ROC-AUC >= 0.75 before container build.\n"
                    "3. Containerized FastAPI serving application receives inference payloads via REST.\n"
                    "4. Statistical Drift Detector computes Kolmogorov-Smirnov (p < 0.05) & PSI (>= 0.20) against reference baseline.\n"
                    "5. Prometheus server scrapes /metrics (OpenMetrics format) every 5s; Grafana visualizes real-time distribution charts.\n"
                    "6. Kubernetes (Kind) cluster orchestrates deployment with Horizontal Pod Autoscaler (HPA 1-5 pods)."
                )
                add_callout(doc, diagram_summary, title="SYSTEM ARCHITECTURE SPECIFICATION")
            else:
                mermaid_buffer.append(line)
            i += 1
            continue

        # Handle generic code blocks
        if line.startswith("```"):
            code_buffer = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_buffer.append(lines[i])
                i += 1
            i += 1  # skip ending ```
            code_text = "\n".join(code_buffer)
            add_callout(doc, code_text, title="CODE / TERMINAL OUTPUT")
            continue

        # Headings
        if line.startswith("## "):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(h_text)
            r.font.name = "Calibri"
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = COLOR_NAVY
            i += 1
            continue

        if line.startswith("### "):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(h_text)
            r.font.name = "Calibri"
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = COLOR_ROYAL
            i += 1
            continue

        if line.startswith("#### "):
            h_text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(h_text)
            r.font.name = "Calibri"
            r.font.size = Pt(11.5)
            r.font.bold = True
            r.font.color.rgb = COLOR_SLATE
            i += 1
            continue

        # Markdown Tables
        if line.startswith("|") and line.endswith("|"):
            table_lines = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            # Filter separator row (e.g., |---|---|)
            parsed_rows = []
            for r_idx, t_line in enumerate(table_lines):
                cells = [c.strip() for c in t_line.strip("|").split("|")]
                if all(re.match(r'^:?-+:?$', c) for c in cells):
                    continue
                parsed_rows.append(cells)

            if parsed_rows:
                num_cols = max(len(r) for r in parsed_rows)
                tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

                for r_idx, row_data in enumerate(parsed_rows):
                    for c_idx in range(num_cols):
                        val = row_data[c_idx] if c_idx < len(row_data) else ""
                        cell = tbl.cell(r_idx, c_idx)
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)

                        # Clean markdown formatting inside cell
                        clean_val = val.replace('$', '').replace('\\text{', '').replace('}', '').replace('\\ge', '>=')
                        r = p.add_run(clean_val)
                        r.font.name = 'Calibri'

                        if r_idx == 0:
                            # Header styling
                            set_cell_background(cell, HEX_TABLE_HEADER)
                            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                            r.font.bold = True
                            r.font.size = Pt(9.5)
                            r.font.color.rgb = RGBColor(255, 255, 255)
                        else:
                            # Body styling
                            if r_idx % 2 == 1:
                                set_cell_background(cell, HEX_ZEBRA)
                            else:
                                set_cell_background(cell, "FFFFFF")
                            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                            r.font.size = Pt(9.0)
                            r.font.color.rgb = COLOR_BODY

                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue

        # Math formulas in block format ($$ ... $$)
        if line.startswith("$$") and line.endswith("$$") and len(line) > 4:
            math_text = line[2:-2].strip()
            add_callout(doc, f"Formula: {math_text}", title="MATHEMATICAL SPECIFICATION")
            i += 1
            continue

        # Bullet List Items
        if line.startswith("- "):
            item_text = line[2:].strip()
            add_styled_paragraph(doc, item_text, style='List Bullet', space_after=3)
            i += 1
            continue

        # Numbered List Items
        if re.match(r'^\d+\.\s', line):
            item_text = re.sub(r'^\d+\.\s', '', line).strip()
            add_styled_paragraph(doc, item_text, style='List Number', space_after=3)
            i += 1
            continue

        # Regular Body Paragraph
        if line:
            # Clean single dollar math markers for document readability
            clean_line = line.replace('$', '')
            add_styled_paragraph(doc, clean_line, style='Normal', space_after=6)

        i += 1

    # Save to both repository and user Downloads folder
    doc.save(str(OUTPUT_DOCX_REPO))
    print(f"Saved repository docx: {OUTPUT_DOCX_REPO}")

    try:
        doc.save(str(OUTPUT_DOCX_DOWNLOADS))
        print(f"Saved Downloads docx: {OUTPUT_DOCX_DOWNLOADS}")
    except Exception as exc:
        print(f"Warning saving to downloads: {exc}")


if __name__ == "__main__":
    main()
